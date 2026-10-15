from fastapi import APIRouter, Depends, HTTPException, Query, status, UploadFile, File
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import or_, desc
from typing import List, Optional
import uuid
import datetime
import os
import time
from io import BytesIO
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

UPLOAD_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), "uploads")
os.makedirs(UPLOAD_DIR, exist_ok=True)


from core.database import get_db
from models.recruitment import VacancyAnnouncement, JobApplication
from schemas.recruitment import (
    VacancyAnnouncementSchema,
    VacancyAnnouncementCreate,
    VacancyAnnouncementUpdate,
    JobApplicationSchema,
    JobApplicationCreate,
    JobApplicationUpdate,
    ApplicationVerificationUpdate
)

router = APIRouter(
    prefix="/api/v1/recruitment",
    tags=["Recruitment & Job Vacancy Portal"]
)

def is_valid_uuid(val: str) -> bool:
    try:
        uuid.UUID(str(val))
        return True
    except (ValueError, TypeError, AttributeError):
        return False


# -------------------------------------------------------------
# 1. SUMMARY STATS (STATIC SUB-ROUTE)
# -------------------------------------------------------------
@router.get("/summary")
def get_recruitment_summary(db: Session = Depends(get_db)):
    total_vacancies = db.query(VacancyAnnouncement).count()
    active_vacancies = db.query(VacancyAnnouncement).filter(VacancyAnnouncement.status == "खुला (Active)").count()
    
    total_applications = db.query(JobApplication).count()
    pending_applications = db.query(JobApplication).filter(JobApplication.status.like("%विचाराधीन%")).count()
    approved_applications = db.query(JobApplication).filter(JobApplication.status.like("%स्वीकृत%")).count()
    rejected_applications = db.query(JobApplication).filter(JobApplication.status.like("%अस्वीकृत%")).count()

    total_openings = db.query(VacancyAnnouncement).with_entities(VacancyAnnouncement.openings_count).all()
    sum_openings = sum([o[0] for o in total_openings if o[0]])

    return {
        "total_vacancies": total_vacancies,
        "active_vacancies": active_vacancies,
        "total_openings": sum_openings,
        "total_applications": total_applications,
        "pending_applications": pending_applications,
        "approved_applications": approved_applications,
        "rejected_applications": rejected_applications
    }


# -------------------------------------------------------------
# 2. EXCEL EXPORT (STATIC SUB-ROUTE)
# -------------------------------------------------------------
@router.get("/export-excel")
def export_recruitment_excel(db: Session = Depends(get_db)):
    vacancies = db.query(VacancyAnnouncement).order_by(desc(VacancyAnnouncement.created_at)).all()
    applications = db.query(JobApplication).order_by(desc(JobApplication.created_at)).all()

    wb = openpyxl.Workbook()

    # SHEET 1: Applications
    ws1 = wb.active
    ws1.title = "दरखास्त_फारम_विवरण"
    ws1.merge_cells("A1:M1")
    t1 = ws1["A1"]
    t1.value = "पाँचपोखरी थाङपाल गाउँपालिका — अनलाइन दरखास्त तथा कर्मचारी पदपूर्ति विवरण"
    t1.font = Font(name="Kalimati", size=13, bold=True, color="FFFFFF")
    t1.fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    t1.alignment = Alignment(horizontal="center", vertical="center")

    headers1 = [
        "क्र.सं.", "दरखास्त नं", "विज्ञापन नं", "पद", "उम्मेदवारको नाम",
        "लिङ्ग", "सम्पर्क फोन", "नागरिकता नं", "ठेगाना", "न्यूनतम योग्यता",
        "अवस्था (Status)", "रोल नम्बर", "परीक्षा केन्द्र"
    ]
    ws1.append([])
    ws1.append(headers1)

    h_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    for col in range(1, len(headers1) + 1):
        c = ws1.cell(row=3, column=col)
        c.font = Font(name="Kalimati", size=10, bold=True, color="FFFFFF")
        c.fill = h_fill
        c.alignment = Alignment(horizontal="center", vertical="center")

    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    for idx, a in enumerate(applications, start=1):
        row = [
            idx,
            a.application_no,
            a.vacancy_no,
            a.post_title,
            a.full_name_np,
            a.gender or "पुरुष",
            a.contact_phone,
            f"{a.citizenship_no or ''} ({a.citizenship_district or ''})",
            a.permanent_address,
            f"{a.qualification_title or ''} - {a.board_university or ''}",
            a.status or "विचाराधीन (Pending)",
            a.admit_card_roll_no or "-",
            a.exam_center or "-"
        ]
        ws1.append(row)
        curr_row = ws1.max_row
        for col_idx in range(1, len(row) + 1):
            cell = ws1.cell(row=curr_row, column=col_idx)
            cell.border = thin_border
            cell.font = Font(name="Kalimati", size=9)

    for col in ws1.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = max(max_len + 3, 13)

    # SHEET 2: Vacancies
    ws2 = wb.create_sheet(title="विज्ञापनहरू")
    ws2.merge_cells("A1:I1")
    t2 = ws2["A1"]
    t2.value = "खुला विज्ञापन तथा पदपूर्ति माग सूची"
    t2.font = Font(name="Kalimati", size=13, bold=True, color="FFFFFF")
    t2.fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    t2.alignment = Alignment(horizontal="center", vertical="center")

    headers2 = [
        "क्र.सं.", "विज्ञापन नं", "पद", "सेवा/समूह", "तह/श्रेणी",
        "माग पद संख्या", "रोजगार प्रकार", "दरखास्त शुल्क (रु.)", "अन्तिम मिति"
    ]
    ws2.append([])
    ws2.append(headers2)

    for col in range(1, len(headers2) + 1):
        c = ws2.cell(row=3, column=col)
        c.font = Font(name="Kalimati", size=10, bold=True, color="FFFFFF")
        c.fill = h_fill
        c.alignment = Alignment(horizontal="center", vertical="center")

    for idx, v in enumerate(vacancies, start=1):
        row2 = [
            idx,
            v.vacancy_no,
            v.title,
            v.service_group,
            v.level,
            v.openings_count,
            v.employment_type,
            float(v.application_fee or 0),
            v.deadline_date_bs
        ]
        ws2.append(row2)
        curr_row = ws2.max_row
        for col_idx in range(1, len(row2) + 1):
            cell = ws2.cell(row=curr_row, column=col_idx)
            cell.border = thin_border
            cell.font = Font(name="Kalimati", size=9)

    for col in ws2.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws2.column_dimensions[col_letter].width = max(max_len + 3, 13)

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)

    filename = f"Paanchpokhari_eRecruitment_Report_{datetime.datetime.now().strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


# -------------------------------------------------------------
# 3. DIGITAL ADMIT CARD LOOKUP (STATIC SUB-ROUTE)
# -------------------------------------------------------------
@router.get("/admit-card")
def get_admit_card(
    application_no: Optional[str] = None,
    phone_or_citizenship: Optional[str] = None,
    db: Session = Depends(get_db)
):
    query = db.query(JobApplication)
    if application_no:
        app = query.filter(JobApplication.application_no == application_no.strip()).first()
    elif phone_or_citizenship:
        q_val = phone_or_citizenship.strip()
        app = query.filter(
            or_(
                JobApplication.contact_phone == q_val,
                JobApplication.citizenship_no == q_val,
                JobApplication.application_no == q_val
            )
        ).first()
    else:
        raise HTTPException(status_code=400, detail="दरखास्त नं वा नागरिकता/फोन नम्बर प्रदान गर्नुहोस्।")

    if not app:
        raise HTTPException(status_code=404, detail="उल्लेखित विवरण अनुसार कुनै दरखास्त फेला परेन।")

    vacancy = db.query(VacancyAnnouncement).filter(
        or_(
            VacancyAnnouncement.id == app.vacancy_id,
            VacancyAnnouncement.vacancy_no == app.vacancy_no
        )
    ).first()

    return {
        "application": JobApplicationSchema.from_orm(app),
        "vacancy": VacancyAnnouncementSchema.from_orm(vacancy) if vacancy else None,
        "is_admit_card_ready": app.status == "स्वीकृत (Approved)" or bool(app.admit_card_roll_no)
    }


# -------------------------------------------------------------
# 3.1 DOCUMENT UPLOAD ENDPOINT
# -------------------------------------------------------------
@router.post("/upload-document")
async def upload_recruitment_document(file: UploadFile = File(...)):
    raw_name = file.filename or "document.jpg"
    _, ext = os.path.splitext(raw_name)
    ext_lower = ext.lower()
    
    allowed = {".jpg", ".jpeg", ".png", ".webp", ".pdf", ".svg", ".doc", ".docx"}
    if ext_lower not in allowed:
        raise HTTPException(
            status_code=400,
            detail="अमान्य फाइल ढाँचा! कृपया तस्बिर (.jpg, .jpeg, .png) वा (.pdf) फाइल मात्र छनोट गर्नुहोस्।"
        )
        
    contents = await file.read()
    if len(contents) > 15 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="फाइलको आकार १५ MB भन्दा बढी हुन सक्दैन।")
        
    unique_suffix = f"{int(time.time() * 1000)}-{os.urandom(4).hex()}"
    filename = f"recruitment-{unique_suffix}{ext_lower}"
    file_path = os.path.join(UPLOAD_DIR, filename)
    
    with open(file_path, "wb") as f:
        f.write(contents)
        
    return {
        "url": f"/uploads/{filename}",
        "filename": filename,
        "original_name": raw_name,
        "size": len(contents),
        "content_type": file.content_type
    }



# -------------------------------------------------------------
# 4. VACANCIES CRUD
# -------------------------------------------------------------
@router.get("/vacancies", response_model=List[VacancyAnnouncementSchema])
def list_vacancies(
    status: Optional[str] = None,
    search: Optional[str] = None,
    db: Session = Depends(get_db)
):
    query = db.query(VacancyAnnouncement)
    if status and status != "all":
        query = query.filter(VacancyAnnouncement.status == status)
    if search:
        s = f"%{search}%"
        query = query.filter(
            or_(
                VacancyAnnouncement.title.ilike(s),
                VacancyAnnouncement.vacancy_no.ilike(s),
                VacancyAnnouncement.service_group.ilike(s)
            )
        )
    return query.order_by(desc(VacancyAnnouncement.created_at)).all()


@router.post("/vacancies", response_model=VacancyAnnouncementSchema)
def create_vacancy(
    data: VacancyAnnouncementCreate,
    db: Session = Depends(get_db)
):
    # Auto-generate vacancy_no if empty
    if not data.vacancy_no:
        count = db.query(VacancyAnnouncement).count() + 1
        data.vacancy_no = f"०{count}/२०८१/०८२"

    new_vac = VacancyAnnouncement(
        id=str(uuid.uuid4()),
        vacancy_no=data.vacancy_no,
        title=data.title,
        service_group=data.service_group,
        level=data.level,
        openings_count=data.openings_count or 1,
        employment_type=data.employment_type or "करार सेवा (Contract)",
        application_fee=data.application_fee or 500.0,
        min_qualification=data.min_qualification,
        publish_date_bs=data.publish_date_bs,
        deadline_date_bs=data.deadline_date_bs,
        exam_date_bs=data.exam_date_bs,
        status=data.status or "खुला (Active)",
        description=data.description
    )
    db.add(new_vac)
    db.commit()
    db.refresh(new_vac)
    return new_vac


@router.get("/vacancies/{vacancy_id}", response_model=VacancyAnnouncementSchema)
def get_vacancy(vacancy_id: str, db: Session = Depends(get_db)):
    cond = [VacancyAnnouncement.vacancy_no == vacancy_id]
    if is_valid_uuid(vacancy_id) or len(vacancy_id) > 10:
        cond.append(VacancyAnnouncement.id == vacancy_id)

    vac = db.query(VacancyAnnouncement).filter(or_(*cond)).first()
    if not vac:
        raise HTTPException(status_code=404, detail="विज्ञापन फेला परेन।")
    return vac


@router.put("/vacancies/{vacancy_id}", response_model=VacancyAnnouncementSchema)
def update_vacancy(
    vacancy_id: str,
    data: VacancyAnnouncementUpdate,
    db: Session = Depends(get_db)
):
    cond = [VacancyAnnouncement.vacancy_no == vacancy_id]
    if is_valid_uuid(vacancy_id) or len(vacancy_id) > 10:
        cond.append(VacancyAnnouncement.id == vacancy_id)

    vac = db.query(VacancyAnnouncement).filter(or_(*cond)).first()
    if not vac:
        raise HTTPException(status_code=404, detail="विज्ञापन फेला परेन।")

    update_dict = data.dict(exclude_unset=True)
    for key, value in update_dict.items():
        if value is not None and hasattr(vac, key):
            setattr(vac, key, value)

    db.commit()
    db.refresh(vac)
    return vac


@router.delete("/vacancies/{vacancy_id}")
def delete_vacancy(vacancy_id: str, db: Session = Depends(get_db)):
    cond = [VacancyAnnouncement.vacancy_no == vacancy_id]
    if is_valid_uuid(vacancy_id) or len(vacancy_id) > 10:
        cond.append(VacancyAnnouncement.id == vacancy_id)

    vac = db.query(VacancyAnnouncement).filter(or_(*cond)).first()
    if not vac:
        raise HTTPException(status_code=404, detail="विज्ञापन फेला परेन।")

    db.delete(vac)
    db.commit()
    return {"message": "विज्ञापन सफलतापूर्वक हटाइयो।", "success": True}



# -------------------------------------------------------------
# 5. APPLICATIONS CRUD & APPLY
# -------------------------------------------------------------
@router.get("/applications", response_model=List[JobApplicationSchema])
def list_applications(
    vacancy_id: Optional[str] = None,
    status: Optional[str] = None,
    search: Optional[str] = None,
    db: Session = Depends(get_db)
):
    query = db.query(JobApplication)
    if vacancy_id and vacancy_id != "all":
        query = query.filter(
            or_(
                JobApplication.vacancy_id == vacancy_id,
                JobApplication.vacancy_no == vacancy_id
            )
        )
    if status and status != "all":
        query = query.filter(JobApplication.status.like(f"%{status}%"))
    if search:
        s = f"%{search}%"
        query = query.filter(
            or_(
                JobApplication.full_name_np.ilike(s),
                JobApplication.application_no.ilike(s),
                JobApplication.contact_phone.ilike(s),
                JobApplication.citizenship_no.ilike(s),
                JobApplication.post_title.ilike(s)
            )
        )
    return query.order_by(desc(JobApplication.created_at)).all()


@router.post("/apply", response_model=JobApplicationSchema)
def submit_job_application(
    data: JobApplicationCreate,
    db: Session = Depends(get_db)
):
    # Lookup vacancy
    vac = None
    if data.vacancy_id:
        cond = [VacancyAnnouncement.id == data.vacancy_id, VacancyAnnouncement.vacancy_no == data.vacancy_id]
        vac = db.query(VacancyAnnouncement).filter(or_(*cond)).first()

    vacancy_no = data.vacancy_no or (vac.vacancy_no if vac else "२०८१/०८२")
    post_title = data.post_title or (vac.title if vac else "पद")

    # Generate sequential application number: APP-2081-XXXX
    count = db.query(JobApplication).count() + 1
    application_no = f"APP-2081-{count:04d}"

    new_app = JobApplication(
        id=str(uuid.uuid4()),
        application_no=application_no,
        vacancy_id=vac.id if vac else data.vacancy_id,
        vacancy_no=vacancy_no,
        post_title=post_title,
        full_name_np=data.full_name_np,
        full_name_en=data.full_name_en,
        dob_bs=data.dob_bs,
        gender=data.gender or "पुरुष",
        citizenship_no=data.citizenship_no,
        citizenship_district=data.citizenship_district,
        contact_phone=data.contact_phone,
        email=data.email,
        permanent_address=data.permanent_address,
        temporary_address=data.temporary_address,
        qualification_title=data.qualification_title,
        board_university=data.board_university,
        passed_year_bs=data.passed_year_bs,
        division_gpa=data.division_gpa,
        council_reg_no=data.council_reg_no,
        experience_years=data.experience_years or 0,
        payment_method=data.payment_method or "eSewa / Khalti / ConnectIPS",
        payment_voucher_no=data.payment_voucher_no or f"TXN-{uuid.uuid4().hex[:8].upper()}",
        photo_url=data.photo_url,
        signature_url=data.signature_url,
        citizenship_doc_url=data.citizenship_doc_url,
        qualification_doc_url=data.qualification_doc_url,
        council_doc_url=data.council_doc_url,
        experience_doc_url=data.experience_doc_url,
        payment_voucher_doc_url=data.payment_voucher_doc_url,
        documents_json=data.documents_json,
        status="विचाराधीन (Pending)"
    )
    db.add(new_app)
    db.commit()
    db.refresh(new_app)
    return new_app


@router.put("/applications/{application_id}/verify", response_model=JobApplicationSchema)
def verify_job_application(
    application_id: str,
    data: ApplicationVerificationUpdate,
    db: Session = Depends(get_db)
):
    cond = [JobApplication.application_no == application_id]
    if is_valid_uuid(application_id) or len(application_id) > 10:
        cond.append(JobApplication.id == application_id)

    app = db.query(JobApplication).filter(or_(*cond)).first()
    if not app:
        raise HTTPException(status_code=404, detail="दरखास्त फेला परेन।")

    app.status = data.status
    if data.rejection_reason is not None:
        app.rejection_reason = data.rejection_reason

    # If approved, ensure admit card details are populated
    if "स्वीकृत" in data.status:
        if data.admit_card_roll_no:
            app.admit_card_roll_no = data.admit_card_roll_no
        elif not app.admit_card_roll_no:
            count = db.query(JobApplication).filter(JobApplication.admit_card_roll_no.isnot(None)).count() + 1
            app.admit_card_roll_no = f"2081-ROLL-{count:03d}"

        app.exam_center = data.exam_center or app.exam_center or "पाँचपोखरी थाङपाल मा.वि., वडा नं ६, सिन्धुपाल्चोक"
        app.exam_date_time = data.exam_date_time or app.exam_date_time or "२०८१/०८/०५ बिहान ११:०० बजे"
    elif "अस्वीकृत" in data.status:
        app.admit_card_roll_no = None

    db.commit()
    db.refresh(app)
    return app


@router.put("/applications/{application_id}", response_model=JobApplicationSchema)
def update_job_application(
    application_id: str,
    data: JobApplicationUpdate,
    db: Session = Depends(get_db)
):
    cond = [JobApplication.application_no == application_id]
    if is_valid_uuid(application_id) or len(application_id) > 10:
        cond.append(JobApplication.id == application_id)

    app = db.query(JobApplication).filter(or_(*cond)).first()
    if not app:
        raise HTTPException(status_code=404, detail="दरखास्त फेला परेन।")

    update_dict = data.dict(exclude_unset=True)
    for key, value in update_dict.items():
        if value is not None and hasattr(app, key):
            setattr(app, key, value)

    db.commit()
    db.refresh(app)
    return app


@router.delete("/applications/{application_id}")
def delete_job_application(application_id: str, db: Session = Depends(get_db)):
    cond = [JobApplication.application_no == application_id]
    if is_valid_uuid(application_id) or len(application_id) > 10:
        cond.append(JobApplication.id == application_id)

    app = db.query(JobApplication).filter(or_(*cond)).first()
    if not app:
        raise HTTPException(status_code=404, detail="दरखास्त फेला परेन।")

    db.delete(app)
    db.commit()
    return {"message": "दरखास्त आवेदन सफलतापूर्वक हटाइयो।", "success": True}

