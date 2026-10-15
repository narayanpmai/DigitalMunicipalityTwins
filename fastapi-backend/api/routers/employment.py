import uuid
import datetime
from io import BytesIO
from typing import List, Optional
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from core.database import get_db
from models.employment import (
    EmploymentLabourBank,
    EmploymentProject,
    EmploymentWorkLog,
    EmploymentLabourDemand
)
from schemas.employment import (
    EmploymentLabourBankSchema,
    EmploymentLabourBankCreate,
    EmploymentLabourBankUpdate,
    ProjectAllocationRequest,
    EmploymentProjectSchema,
    EmploymentProjectCreate,
    EmploymentWorkLogSchema,
    WorkLogCreate,
    EmploymentLabourDemandSchema,
    EmploymentLabourDemandCreate
)

router = APIRouter(
    prefix="/api/v1/employment",
    tags=["Employment Service Centre & Labour Bank (रोजगार सेवा केन्द्र तथा श्रम बैंक)"]
)


@router.get('/job-seekers')
def get_employment_job_seekers():
    return []


@router.get('/stats')
def get_employment_stats():
    return {'total': 0}

# -------------------------------------------------------------
# HELPERS
# -------------------------------------------------------------
def is_valid_uuid(val: str) -> bool:
    try:
        uuid.UUID(str(val))
        return True
    except (ValueError, AttributeError):
        return False

def find_worker(db: Session, id_or_worker_id: str) -> Optional[EmploymentLabourBank]:
    if is_valid_uuid(id_or_worker_id):
        return db.query(EmploymentLabourBank).filter(
            (EmploymentLabourBank.id == id_or_worker_id) | (EmploymentLabourBank.worker_id == id_or_worker_id)
        ).first()
    return db.query(EmploymentLabourBank).filter(
        EmploymentLabourBank.worker_id == id_or_worker_id
    ).first()

def serialize_worker(w: EmploymentLabourBank) -> dict:
    daily_wage = float(w.daily_wage_rate or 850.0)
    total_wage = float(w.total_wage_paid or 0.0)
    days_worked = int(w.total_days_worked or 0)
    allocated_days = int(w.guaranteed_days_allocated or 100)
    
    return {
        "id": str(w.id),
        "workerId": w.worker_id or "",
        "worker_id": w.worker_id or "",
        "fullName": w.full_name or "",
        "full_name": w.full_name or "",
        "citizenshipNo": w.citizenship_no or "",
        "citizenship_no": w.citizenship_no or "",
        "gender": w.gender or "पुरुष",
        "wardNumber": int(w.ward_number or 1),
        "ward_number": int(w.ward_number or 1),
        "tole": w.tole_village or "",
        "toleVillage": w.tole_village or "",
        "tole_village": w.tole_village or "",
        "phone": w.contact_phone or "",
        "contactPhone": w.contact_phone or "",
        "contact_phone": w.contact_phone or "",
        "skillCategory": w.skill_category or "सामान्य निर्माण श्रमिक (General Construction)",
        "skill_category": w.skill_category or "सामान्य निर्माण श्रमिक (General Construction)",
        "skillLevel": w.skill_level or "सामान्य (Unskilled)",
        "skill_level": w.skill_level or "सामान्य (Unskilled)",
        "dailyWageRate": daily_wage,
        "daily_wage_rate": daily_wage,
        "guaranteedDaysAllocated": allocated_days,
        "guaranteed_days_allocated": allocated_days,
        "totalDaysWorked": days_worked,
        "total_days_worked": days_worked,
        "totalWagePaid": total_wage,
        "total_wage_paid": total_wage,
        "allocatedProject": w.allocated_project or "",
        "allocated_project": w.allocated_project or "",
        "bankName": w.bank_name or "",
        "bank_name": w.bank_name or "",
        "bankAccountNo": w.bank_account_no or "",
        "bank_account_no": w.bank_account_no or "",
        "status": w.status or "सूचीकृत (Registered)",
        "pmepFiscalYear": w.pmep_fiscal_year or "२०८१/०८२",
        "pmep_fiscal_year": w.pmep_fiscal_year or "२०८१/०८२",
        "registeredDateBs": w.registered_date_bs or "",
        "registered_date_bs": w.registered_date_bs or "",
        "notes": w.notes or ""
    }


# -------------------------------------------------------------
# 1. SUMMARY STATS (MUST BE FIRST)
# -------------------------------------------------------------
@router.get("/summary")
def get_employment_summary(db: Session = Depends(get_db)):
    workers = db.query(EmploymentLabourBank).all()
    total_registered = len(workers)
    
    total_deployed = sum(1 for w in workers if w.allocated_project or 'परिचालित' in (w.status or ''))
    completed_100 = sum(1 for w in workers if (w.total_days_worked or 0) >= 100)
    total_person_days = sum(int(w.total_days_worked or 0) for w in workers)
    total_wage_disbursed = sum(float(w.total_wage_paid or 0.0) for w in workers)
    
    # Skill breakdown
    skill_map = {}
    for w in workers:
        sk = w.skill_category or 'सामान्य निर्माण श्रमिक'
        skill_map[sk] = skill_map.get(sk, 0) + 1
    skill_breakdown = [{"skill": k, "count": str(v)} for k, v in skill_map.items()]
    
    # Ward breakdown
    ward_map = {w: 0 for w in range(1, 9)}
    for w in workers:
        wn = w.ward_number or 1
        ward_map[wn] = ward_map.get(wn, 0) + 1
    ward_breakdown = [{"ward": k, "count": str(v)} for k, v in ward_map.items()]
    
    # Project statistics
    projects_count = db.query(EmploymentProject).count()
    active_projects = db.query(EmploymentProject).filter(EmploymentProject.status.ilike('%सञ्चालनमा%')).count()
    completed_projects = db.query(EmploymentProject).filter(EmploymentProject.status.ilike('%सम्पन्न%')).count()
    total_budget = float(db.query(func.sum(EmploymentProject.allocated_budget)).scalar() or 0.0)

    # Demands count
    total_demands = db.query(EmploymentLabourDemand).count()

    return {
        "totalRegistered": total_registered,
        "totalDeployed": total_deployed,
        "guaranteed100DaysCompleted": completed_100,
        "totalPersonDaysGenerated": total_person_days,
        "totalWageDisbursedNPR": total_wage_disbursed,
        "totalProjects": projects_count,
        "activeProjects": active_projects,
        "completedProjects": completed_projects,
        "totalProjectBudget": total_budget,
        "totalDemands": total_demands,
        "skillBreakdown": skill_breakdown,
        "wardBreakdown": ward_breakdown
    }


# -------------------------------------------------------------
# 2. PMEP PROJECTS (STATIC SUB-ROUTE)
# -------------------------------------------------------------
@router.get("/projects", response_model=List[EmploymentProjectSchema])
@router.get("/projects/", response_model=List[EmploymentProjectSchema])
def get_projects(
    ward: Optional[int] = Query(None),
    status: Optional[str] = Query(None),
    db: Session = Depends(get_db)
):
    q = db.query(EmploymentProject)
    if ward:
        q = q.filter(EmploymentProject.ward_number == ward)
    if status and status != 'all':
        q = q.filter(EmploymentProject.status.ilike(f"%{status}%"))
    return q.order_by(EmploymentProject.ward_number, EmploymentProject.created_at.desc()).all()


@router.post("/projects", response_model=EmploymentProjectSchema)
@router.post("/projects/", response_model=EmploymentProjectSchema)
def create_project(data: EmploymentProjectCreate, db: Session = Depends(get_db)):
    code = data.project_code
    if not code:
        count = db.query(EmploymentProject).count() + 1
        code = f"PMEP-PRJ-081-{count:02d}"
        
    proj = EmploymentProject(
        id=str(uuid.uuid4()),
        project_code=code,
        title=data.title,
        ward_number=data.ward_number,
        tole_village=data.tole_village,
        allocated_budget=data.allocated_budget or 0.0,
        target_person_days=data.target_person_days or 100,
        completed_person_days=0,
        daily_wage_rate=data.daily_wage_rate or 850.0,
        workers_count=data.workers_count or 0,
        supervisor_name=data.supervisor_name,
        supervisor_phone=data.supervisor_phone,
        status=data.status or "सञ्चालनमा (In Progress)",
        fiscal_year=data.fiscal_year or "२०८१/०८२",
        start_date_bs=data.start_date_bs or "२०८१-०२-०१",
        end_date_bs=data.end_date_bs or "२०८१-०५-३०"
    )
    db.add(proj)
    db.commit()
    db.refresh(proj)
    return proj


@router.put("/projects/{id}", response_model=EmploymentProjectSchema)
def update_project(id: str, data: EmploymentProjectCreate, db: Session = Depends(get_db)):
    proj = db.query(EmploymentProject).filter(
        (EmploymentProject.id == id) | (EmploymentProject.project_code == id)
    ).first()
    if not proj:
        raise HTTPException(status_code=404, detail="आयोजना फेला परेन")
        
    for k, v in data.dict(exclude_unset=True).items():
        if hasattr(proj, k) and v is not None:
            setattr(proj, k, v)
            
    db.commit()
    db.refresh(proj)
    return proj


@router.delete("/projects/{id}")
def delete_project(id: str, db: Session = Depends(get_db)):
    proj = db.query(EmploymentProject).filter(
        (EmploymentProject.id == id) | (EmploymentProject.project_code == id)
    ).first()
    if not proj:
        raise HTTPException(status_code=404, detail="आयोजना फेला परेन")
    db.delete(proj)
    db.commit()
    return {"message": "आयोजना सफलतापूर्वक हटाइयो"}


# -------------------------------------------------------------
# 3. WORK LOGS & ATTENDANCE (STATIC SUB-ROUTE)
# -------------------------------------------------------------
@router.get("/work-logs", response_model=List[EmploymentWorkLogSchema])
@router.get("/work-logs/", response_model=List[EmploymentWorkLogSchema])
def get_work_logs(
    worker_id: Optional[str] = Query(None),
    payout_status: Optional[str] = Query(None),
    db: Session = Depends(get_db)
):
    q = db.query(EmploymentWorkLog)
    if worker_id:
        q = q.filter(EmploymentWorkLog.worker_id == worker_id)
    if payout_status and payout_status != 'all':
        q = q.filter(EmploymentWorkLog.payout_status.ilike(f"%{payout_status}%"))
    return q.order_by(EmploymentWorkLog.created_at.desc()).all()


@router.post("/work-logs", response_model=EmploymentWorkLogSchema)
@router.post("/work-logs/", response_model=EmploymentWorkLogSchema)
def create_work_log(data: WorkLogCreate, db: Session = Depends(get_db)):
    wage = (data.daily_wage_rate or 850.0) * data.days_worked
    log = EmploymentWorkLog(
        id=str(uuid.uuid4()),
        worker_id=data.worker_id,
        worker_name=data.worker_name or "श्रमिक",
        project_title=data.project_title,
        ward_number=data.ward_number or 1,
        work_date_bs=data.work_date_bs,
        days_worked=data.days_worked,
        daily_wage_rate=data.daily_wage_rate or 850.0,
        wage_amount=wage,
        payout_status=data.payout_status or "भुक्तानी बाँकी (Pending)",
        voucher_no=data.voucher_no,
        remarks=data.remarks
    )
    db.add(log)
    
    # Also update worker cumulative days
    w = find_worker(db, data.worker_id)
    if w:
        w.total_days_worked = (w.total_days_worked or 0) + data.days_worked
        w.total_wage_paid = float(w.total_wage_paid or 0.0) + wage
        if w.total_days_worked >= 100:
            w.status = "१०० दिन सम्पन्न (Completed)"
            
    db.commit()
    db.refresh(log)
    return log


@router.put("/work-logs/{id}/disburse")
def disburse_work_log(
    id: str,
    voucher_no: Optional[str] = Query(None),
    db: Session = Depends(get_db)
):
    log = db.query(EmploymentWorkLog).filter(EmploymentWorkLog.id == id).first()
    if not log:
        raise HTTPException(status_code=404, detail="कार्य लग फेला परेन")
        
    vch = voucher_no or f"VCH-PMEP-081-{log.ward_number:02d}-{int(datetime.datetime.utcnow().timestamp()) % 10000:04d}"
    log.payout_status = "बैंक निकासा सम्पन्न (Disbursed)"
    log.voucher_no = vch
    db.commit()
    db.refresh(log)
    return log


# -------------------------------------------------------------
# 4. PRIVATE / LOCAL LABOUR DEMANDS (STATIC SUB-ROUTE)
# -------------------------------------------------------------
@router.get("/demands", response_model=List[EmploymentLabourDemandSchema])
def get_demands(db: Session = Depends(get_db)):
    return db.query(EmploymentLabourDemand).order_by(EmploymentLabourDemand.created_at.desc()).all()


@router.post("/demands", response_model=EmploymentLabourDemandSchema)
def create_demand(data: EmploymentLabourDemandCreate, db: Session = Depends(get_db)):
    dem = EmploymentLabourDemand(
        id=str(uuid.uuid4()),
        employer_name=data.employer_name,
        contact_person=data.contact_person,
        contact_phone=data.contact_phone,
        ward_number=data.ward_number,
        skill_required=data.skill_required,
        required_count=data.required_count or 1,
        daily_wage_offered=data.daily_wage_offered or 1000.0,
        work_duration_days=data.work_duration_days or 30,
        status=data.status or "माग दर्ता (Registered)"
    )
    db.add(dem)
    db.commit()
    db.refresh(dem)
    return dem


@router.put("/demands/{id}/status")
def update_demand_status(id: str, status: str = Query(...), db: Session = Depends(get_db)):
    dem = db.query(EmploymentLabourDemand).filter(EmploymentLabourDemand.id == id).first()
    if not dem:
        raise HTTPException(status_code=404, detail="श्रम माग फेला परेन")
    dem.status = status
    db.commit()
    db.refresh(dem)
    return dem


# -------------------------------------------------------------
# 5. COMPREHENSIVE EXCEL EXPORT (STATIC SUB-ROUTE)
# -------------------------------------------------------------
@router.get("/export-excel")
def export_employment_excel(db: Session = Depends(get_db)):
    workers = db.query(EmploymentLabourBank).order_by(EmploymentLabourBank.ward_number, EmploymentLabourBank.worker_id).all()
    projects = db.query(EmploymentProject).order_by(EmploymentProject.ward_number).all()
    logs = db.query(EmploymentWorkLog).order_by(EmploymentWorkLog.created_at.desc()).all()
    
    wb = openpyxl.Workbook()
    
    # SHEET 1: Workers
    ws1 = wb.active
    ws1.title = "श्रम_बैंक_श्रमिक_लगत"
    
    ws1.merge_cells("A1:N1")
    t1 = ws1["A1"]
    t1.value = "पाँचपोखरी थाङपाल गाउँपालिका — रोजगार सेवा केन्द्र तथा स्थानीय श्रम बैंक"
    t1.font = Font(name="Kalimati", size=13, bold=True, color="FFFFFF")
    t1.fill = PatternFill(start_color="065F46", end_color="065F46", fill_type="solid")
    t1.alignment = Alignment(horizontal="center", vertical="center")
    
    headers1 = [
        "क्र.सं.", "श्रमिक संकेत नं", "श्रमिकको नाम", "नागरिकता नं", "लिङ्ग",
        "वडा नं", "टोल/गाउँ", "सीप वर्ग", "दैनिक ज्याला दर (रु.)", "काम गरेको दिन",
        "कुल भुक्तानी (रु.)", "खटाइएको आयोजना", "बैंक तथा खाता नं", "अवस्था"
    ]
    ws1.append([])
    ws1.append(headers1)
    
    h_fill = PatternFill(start_color="047857", end_color="047857", fill_type="solid")
    for col in range(1, len(headers1) + 1):
        c = ws1.cell(row=3, column=col)
        c.font = Font(name="Kalimati", size=10, bold=True, color="FFFFFF")
        c.fill = h_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        
    for idx, w in enumerate(workers, start=1):
        bank_info = f"{w.bank_name or ''} ({w.bank_account_no or ''})" if w.bank_account_no else "उपलब्ध छैन"
        row = [
            idx,
            w.worker_id or "",
            w.full_name or "",
            w.citizenship_no or "",
            w.gender or "पुरुष",
            w.ward_number or 1,
            w.tole_village or "",
            w.skill_category or "",
            float(w.daily_wage_rate or 850.0),
            int(w.total_days_worked or 0),
            float(w.total_wage_paid or 0.0),
            w.allocated_project or "खटाउन बाँकी",
            bank_info,
            w.status or "सूचीकृत"
        ]
        ws1.append(row)
        
    for col in ws1.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = max(max_len + 3, 12)
        
    # SHEET 2: Projects
    ws2 = wb.create_sheet(title="रोजगारी_आयोजनाहरू")
    ws2.merge_cells("A1:J1")
    p1 = ws2["A1"]
    p1.value = "प्रधानमन्त्री रोजगार कार्यक्रम (PMEP) तथा स्थानीय पूर्वाधार विकास आयोजनाहरू"
    p1.font = Font(name="Kalimati", size=13, bold=True, color="FFFFFF")
    p1.fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    p1.alignment = Alignment(horizontal="center", vertical="center")
    
    headers2 = [
        "क्र.सं.", "आयोजना संकेत", "आयोजनाको नाम", "वडा नं", "विनियोजित बजेट (रु.)",
        "लक्ष्य मानव-दिन", "सम्पन्न मानव-दिन", "परिचालित श्रमिक", "प्राविधिक / सुपरभाइजर", "अवस्था"
    ]
    ws2.append([])
    ws2.append(headers2)
    
    p_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    for col in range(1, len(headers2) + 1):
        c = ws2.cell(row=3, column=col)
        c.font = Font(name="Kalimati", size=10, bold=True, color="FFFFFF")
        c.fill = p_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        
    for idx, p in enumerate(projects, start=1):
        row2 = [
            idx,
            p.project_code,
            p.title,
            p.ward_number,
            float(p.allocated_budget or 0.0),
            p.target_person_days,
            p.completed_person_days,
            p.workers_count,
            f"{p.supervisor_name or ''} ({p.supervisor_phone or ''})",
            p.status
        ]
        ws2.append(row2)
        
    for col in ws2.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws2.column_dimensions[col_letter].width = max(max_len + 3, 14)

    # SHEET 3: Work logs
    ws3 = wb.create_sheet(title="कार्य_दिन_र_ज्याला_लग")
    ws3.merge_cells("A1:J1")
    l1 = ws3["A1"]
    l1.value = "कार्य दिन तथा ज्याला भुक्तानी भरपाई विवरण"
    l1.font = Font(name="Kalimati", size=13, bold=True, color="FFFFFF")
    l1.fill = PatternFill(start_color="B45309", end_color="B45309", fill_type="solid")
    l1.alignment = Alignment(horizontal="center", vertical="center")
    
    headers3 = [
        "क्र.सं.", "श्रमिक संकेत", "श्रमिकको नाम", "आयोजनाको नाम", "कार्य मिति (वि.सं.)",
        "काम गरेको दिन", "दैनिक ज्याला", "ज्याला रकम (रु.)", "भुक्तानी अवस्था", "भरपाई नं"
    ]
    ws3.append([])
    ws3.append(headers3)
    
    l_fill = PatternFill(start_color="D97706", end_color="D97706", fill_type="solid")
    for col in range(1, len(headers3) + 1):
        c = ws3.cell(row=3, column=col)
        c.font = Font(name="Kalimati", size=10, bold=True, color="FFFFFF")
        c.fill = l_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        
    for idx, l in enumerate(logs, start=1):
        row3 = [
            idx,
            l.worker_id,
            l.worker_name,
            l.project_title,
            l.work_date_bs,
            l.days_worked,
            float(l.daily_wage_rate or 850.0),
            float(l.wage_amount or 0.0),
            l.payout_status,
            l.voucher_no or "-"
        ]
        ws3.append(row3)
        
    for col in ws3.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws3.column_dimensions[col_letter].width = max(max_len + 3, 14)
        
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    
    filename = f"Paanchpokhari_ESC_LabourBank_Report_{datetime.datetime.now().strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


# -------------------------------------------------------------
# 6. LABOUR BANK WORKERS (ROOT LIST & DYNAMIC ID ROUTE)
# -------------------------------------------------------------
@router.get("")
@router.get("/")
@router.get("/labour-bank")
@router.get("/labour-bank/")
@router.get("/workers")
def get_workers(
    ward: Optional[int] = Query(None, description="वडा नम्बर"),
    skill: Optional[str] = Query(None, description="सीप वर्ग"),
    status: Optional[str] = Query(None, description="अवस्था"),
    search: Optional[str] = Query(None, description="नाम, संकेत वा फोन"),
    db: Session = Depends(get_db)
):
    query = db.query(EmploymentLabourBank)
    if ward:
        query = query.filter(EmploymentLabourBank.ward_number == ward)
    if skill and skill != 'all':
        query = query.filter(EmploymentLabourBank.skill_category == skill)
    if status and status != 'all':
        query = query.filter(EmploymentLabourBank.status == status)
    if search:
        s = f"%{search}%"
        query = query.filter(
            (EmploymentLabourBank.full_name.ilike(s)) |
            (EmploymentLabourBank.worker_id.ilike(s)) |
            (EmploymentLabourBank.citizenship_no.ilike(s)) |
            (EmploymentLabourBank.contact_phone.ilike(s))
        )
    
    workers = query.order_by(EmploymentLabourBank.ward_number, EmploymentLabourBank.worker_id).all()
    return [serialize_worker(w) for w in workers]


@router.post("")
@router.post("/")
@router.post("/labour-bank")
@router.post("/labour-bank/")
def create_worker(data: EmploymentLabourBankCreate, db: Session = Depends(get_db)):
    w_id = data.worker_id
    if not w_id:
        count = db.query(EmploymentLabourBank).count() + 1
        w_id = f"ESC-2081-W0{data.ward_number}-{count:03d}"
        
    db_worker = EmploymentLabourBank(
        id=str(uuid.uuid4()),
        worker_id=w_id,
        full_name=data.full_name,
        citizenship_no=data.citizenship_no,
        gender=data.gender or "पुरुष",
        ward_number=data.ward_number,
        tole_village=data.tole_village,
        contact_phone=data.contact_phone,
        skill_category=data.skill_category or "सामान्य निर्माण श्रमिक (General Construction)",
        skill_level=data.skill_level or "सामान्य (Unskilled)",
        daily_wage_rate=data.daily_wage_rate or 850.0,
        guaranteed_days_allocated=data.guaranteed_days_allocated or 100,
        total_days_worked=data.total_days_worked or 0,
        total_wage_paid=data.total_wage_paid or 0.0,
        allocated_project=data.allocated_project,
        bank_name=data.bank_name or "नेपाल बैंक लि., भोताङ शाखा",
        bank_account_no=data.bank_account_no,
        status=data.status or "सूचीकृत (Registered)",
        pmep_fiscal_year=data.pmep_fiscal_year or "२०८१/०८२",
        registered_date_bs=data.registered_date_bs or "२०८१-०१-१५",
        notes=data.notes
    )
    db.add(db_worker)
    db.commit()
    db.refresh(db_worker)
    return serialize_worker(db_worker)


@router.get("/{id}")
def get_worker_by_id(id: str, db: Session = Depends(get_db)):
    w = find_worker(db, id)
    if not w:
        raise HTTPException(status_code=404, detail="श्रमिक लगत फेला परेन")
    return serialize_worker(w)


@router.put("/{id}")
def update_worker(id: str, updates: EmploymentLabourBankUpdate, db: Session = Depends(get_db)):
    w = find_worker(db, id)
    if not w:
        raise HTTPException(status_code=404, detail="श्रमिक लगत फेला परेन")
        
    update_data = updates.dict(exclude_unset=True)
    for k, v in update_data.items():
        if hasattr(w, k) and v is not None:
            setattr(w, k, v)
            
    db.commit()
    db.refresh(w)
    return serialize_worker(w)


@router.put("/{id}/allocate")
def allocate_worker_to_project(id: str, req: ProjectAllocationRequest, db: Session = Depends(get_db)):
    w = find_worker(db, id)
    if not w:
        raise HTTPException(status_code=404, detail="श्रमिक लगत फेला परेन")
        
    w.allocated_project = req.projectTitle
    if req.dailyWageRate:
        w.daily_wage_rate = req.dailyWageRate
    if req.days:
        w.guaranteed_days_allocated = req.days
    w.status = "परिचालित (Deployed)"
    
    # Increment project workers_count if project found
    proj = db.query(EmploymentProject).filter(EmploymentProject.title == req.projectTitle).first()
    if proj:
        proj.workers_count = (proj.workers_count or 0) + 1
        
    db.commit()
    db.refresh(w)
    return serialize_worker(w)


@router.put("/{id}/log-work")
def log_worker_days(
    id: str,
    days: int = Query(..., description="थप कार्य दिन संख्या"),
    work_date_bs: str = Query("२०८१-०३-१५", description="कार्य मिति"),
    remarks: Optional[str] = Query(None, description="कैफियत"),
    db: Session = Depends(get_db)
):
    w = find_worker(db, id)
    if not w:
        raise HTTPException(status_code=404, detail="श्रमिक लगत फेला परेन")
        
    prev_days = int(w.total_days_worked or 0)
    w.total_days_worked = prev_days + days
    wage_addition = float(w.daily_wage_rate or 850.0) * days
    w.total_wage_paid = float(w.total_wage_paid or 0.0) + wage_addition
    
    if w.total_days_worked >= 100:
        w.status = "१०० दिन सम्पन्न (Completed)"
    else:
        w.status = "परिचालित (Deployed)"
        
    # Also create Work Log entry
    log = EmploymentWorkLog(
        id=str(uuid.uuid4()),
        worker_id=w.worker_id,
        worker_name=w.full_name,
        project_title=w.allocated_project or "स्थानीय पूर्वाधार मर्मत",
        ward_number=w.ward_number,
        work_date_bs=work_date_bs,
        days_worked=days,
        daily_wage_rate=float(w.daily_wage_rate or 850.0),
        wage_amount=wage_addition,
        payout_status="भुक्तानी बाँकी (Pending)",
        remarks=remarks or "कार्य दिन लग प्रविष्टि"
    )
    db.add(log)
    
    # Update project completed person days if project title matches
    if w.allocated_project:
        proj = db.query(EmploymentProject).filter(EmploymentProject.title == w.allocated_project).first()
        if proj:
            proj.completed_person_days = (proj.completed_person_days or 0) + days
            
    db.commit()
    db.refresh(w)
    return serialize_worker(w)


@router.delete("/{id}")
def delete_worker(id: str, db: Session = Depends(get_db)):
    w = find_worker(db, id)
    if not w:
        raise HTTPException(status_code=404, detail="श्रमिक लगत फेला परेन")
    db.delete(w)
    db.commit()
    return {"message": "श्रमिक लगत सफलतापूर्वक हटाइयो"}
