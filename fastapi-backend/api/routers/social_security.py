import os
import uuid
import datetime
from io import BytesIO
from typing import List, Optional
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from core.database import get_db
from models.social_security import SocialSecurityBeneficiary
from schemas.social_security import (
    SocialSecurityBeneficiarySchema,
    SocialSecurityBeneficiaryCreate,
    SocialSecurityBeneficiaryUpdate
)

router = APIRouter(
    prefix="/api/v1/social-security",
    tags=["Social Security (सामाजिक सुरक्षा)"]
)

# Allowance categories and standard monthly rates (MoFAGA/DoNIDCR standard)
CATEGORY_RATES = {
    'जेष्ठ नागरिक (६८ वर्ष माथि)': 4000.0,
    'जेष्ठ नागरिक दलित (६० वर्ष माथि)': 2660.0,
    'एकल महिला / विधवा (६० वर्ष माथि)': 2660.0,
    'पूर्ण अपाङ्गता (क वर्ग - रातो कार्ड)': 3990.0,
    'अति अशक्त अपाङ्गता (ख वर्ग - नीलो कार्ड)': 2128.0,
    'बालबालिका पोषण भत्ता (५ वर्ष मुनि)': 532.0,
    'लोपोन्मुख आदिवासी / जनजाति': 3990.0,
}

@router.get("", response_model=List[SocialSecurityBeneficiarySchema])
@router.get("/", response_model=List[SocialSecurityBeneficiarySchema])
@router.get("/beneficiaries", response_model=List[SocialSecurityBeneficiarySchema])
@router.get("/beneficiaries/", response_model=List[SocialSecurityBeneficiarySchema])
def get_beneficiaries(
    ward: Optional[int] = Query(None, description="वडा नम्बर"),
    category: Optional[str] = Query(None, description="भत्ताको प्रकार"),
    status: Optional[str] = Query(None, description="अवस्था (सक्रिय, नवीकृत, आदि)"),
    search: Optional[str] = Query(None, description="नाम वा परिचयपत्र नम्बर"),
    db: Session = Depends(get_db)
):
    query = db.query(SocialSecurityBeneficiary)
    if ward:
        query = query.filter(SocialSecurityBeneficiary.ward_number == ward)
    if category and category != 'all':
        query = query.filter(SocialSecurityBeneficiary.category == category)
    if status and status != 'all':
        query = query.filter(SocialSecurityBeneficiary.status == status)
    if search:
        s = f"%{search}%"
        query = query.filter(
            (SocialSecurityBeneficiary.full_name.ilike(s)) |
            (SocialSecurityBeneficiary.beneficiary_id.ilike(s)) |
            (SocialSecurityBeneficiary.citizenship_or_nid_no.ilike(s)) |
            (SocialSecurityBeneficiary.bank_account_no.ilike(s))
        )
    return query.order_by(SocialSecurityBeneficiary.ward_number, SocialSecurityBeneficiary.full_name).all()

@router.get("/summary")
def get_social_security_summary(db: Session = Depends(get_db)):
    beneficiaries = db.query(SocialSecurityBeneficiary).all()
    total_count = len(beneficiaries)
    
    total_monthly = sum(float(b.monthly_allowance_npr or 0) for b in beneficiaries)
    total_quarterly = sum(float(b.quarterly_allowance_npr or (b.monthly_allowance_npr or 0) * 3) for b in beneficiaries)
    
    active_count = sum(1 for b in beneficiaries if b.status in ['सक्रिय', 'नवीकृत'])
    
    # Category breakdown
    cat_counts = {}
    for b in beneficiaries:
        c = b.category or 'अन्य'
        cat_counts[c] = cat_counts.get(c, 0) + 1
        
    # Ward breakdown
    ward_counts = {}
    for w in range(1, 9):
        ward_counts[w] = 0
    for b in beneficiaries:
        w = b.ward_number or 1
        ward_counts[w] = ward_counts.get(w, 0) + 1
        
    return {
        "totalBeneficiaries": total_count,
        "activeBeneficiaries": active_count,
        "totalMonthlyPayout": total_monthly,
        "totalQuarterlyPayout": total_quarterly,
        "byCategory": cat_counts,
        "byWard": ward_counts
    }

@router.post("", response_model=SocialSecurityBeneficiarySchema)
@router.post("/", response_model=SocialSecurityBeneficiarySchema)
@router.post("/beneficiaries", response_model=SocialSecurityBeneficiarySchema)
@router.post("/beneficiaries/", response_model=SocialSecurityBeneficiarySchema)
def create_beneficiary(beneficiary: SocialSecurityBeneficiaryCreate, db: Session = Depends(get_db)):
    # Auto-generate beneficiary_id if empty
    b_id = beneficiary.beneficiary_id
    if not b_id:
        count = db.query(SocialSecurityBeneficiary).count() + 1
        b_id = f"SSIS-2081-W0{beneficiary.ward_number}-{count:04d}"
        
    # Standard rates check
    monthly = beneficiary.monthly_allowance_npr
    if not monthly or monthly == 4000.0:
        monthly = CATEGORY_RATES.get(beneficiary.category, 4000.0)
    quarterly = monthly * 3
    
    db_beneficiary = SocialSecurityBeneficiary(
        id=str(uuid.uuid4()),
        beneficiary_id=b_id,
        full_name=beneficiary.full_name,
        full_name_en=beneficiary.full_name_en,
        citizenship_or_nid_no=beneficiary.citizenship_or_nid_no,
        category=beneficiary.category,
        ward_number=beneficiary.ward_number,
        tole_village=beneficiary.tole_village,
        gender=beneficiary.gender,
        date_of_birth_bs=beneficiary.date_of_birth_bs,
        contact_phone=beneficiary.contact_phone,
        guardian_name=beneficiary.guardian_name,
        bank_name=beneficiary.bank_name,
        bank_account_no=beneficiary.bank_account_no,
        monthly_allowance_npr=monthly,
        quarterly_allowance_npr=quarterly,
        status=beneficiary.status or 'सक्रिय',
        fiscal_year=beneficiary.fiscal_year or '२०८१/०८२',
        renewal_date_bs=beneficiary.renewal_date_bs or '२०८१-०४-०१'
    )
    db.add(db_beneficiary)
    db.commit()
    db.refresh(db_beneficiary)
    return db_beneficiary

@router.put("/beneficiaries/{id}/renew", response_model=SocialSecurityBeneficiarySchema)
@router.put("/{id}/renew", response_model=SocialSecurityBeneficiarySchema)
def renew_beneficiary(id: str, db: Session = Depends(get_db)):
    b = db.query(SocialSecurityBeneficiary).filter(
        (SocialSecurityBeneficiary.id == id) | (SocialSecurityBeneficiary.beneficiary_id == id)
    ).first()
    if not b:
        raise HTTPException(status_code=404, detail="लाभग्राही फेला परेन")
    b.status = "नवीकृत"
    b.renewal_date_bs = "२०८१-०४-०१"
    db.commit()
    db.refresh(b)
    return b

@router.put("/beneficiaries/{id}/cancel")
@router.put("/{id}/cancel")
def cancel_beneficiary(id: str, db: Session = Depends(get_db)):
    b = db.query(SocialSecurityBeneficiary).filter(
        (SocialSecurityBeneficiary.id == id) | (SocialSecurityBeneficiary.beneficiary_id == id)
    ).first()
    if not b:
        raise HTTPException(status_code=404, detail="लाभग्राही फेला परेन")
    b.status = "कट्टा"
    db.commit()
    db.refresh(b)
    return {"message": "लाभग्राही लगत कट्टा गरियो"}

@router.put("/beneficiaries/{id}", response_model=SocialSecurityBeneficiarySchema)
def update_beneficiary(id: str, updates: SocialSecurityBeneficiaryUpdate, db: Session = Depends(get_db)):
    b = db.query(SocialSecurityBeneficiary).filter(
        (SocialSecurityBeneficiary.id == id) | (SocialSecurityBeneficiary.beneficiary_id == id)
    ).first()
    if not b:
        raise HTTPException(status_code=404, detail="लाभग्राही फेला परेन")
        
    update_data = updates.dict(exclude_unset=True)
    for field, val in update_data.items():
        setattr(b, field, val)
        
    if 'monthly_allowance_npr' in update_data and not update_data.get('quarterly_allowance_npr'):
        b.quarterly_allowance_npr = float(b.monthly_allowance_npr or 0) * 3
        
    db.commit()
    db.refresh(b)
    return b

@router.delete("/beneficiaries/{id}")
def delete_beneficiary(id: str, db: Session = Depends(get_db)):
    b = db.query(SocialSecurityBeneficiary).filter(
        (SocialSecurityBeneficiary.id == id) | (SocialSecurityBeneficiary.beneficiary_id == id)
    ).first()
    if not b:
        raise HTTPException(status_code=404, detail="लाभग्राही फेला परेन")
    db.delete(b)
    db.commit()
    return {"message": "लाभग्राही लगत सफलतापूर्वक हटाइयो"}

@router.get("/export-excel")
def export_beneficiaries_excel(db: Session = Depends(get_db)):
    beneficiaries = db.query(SocialSecurityBeneficiary).order_by(
        SocialSecurityBeneficiary.ward_number,
        SocialSecurityBeneficiary.full_name
    ).all()
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "सामाजिक_सुरक्षा_लाभग्राही"
    
    ws.merge_cells("A1:K1")
    t1 = ws["A1"]
    t1.value = "पाँचपोखरी थाङपाल गाउँपालिका - सामाजिक सुरक्षा तथा व्यक्तिगत घटना दर्ता प्रणाली"
    t1.font = Font(name="Kalimati", size=14, bold=True, color="1E3A8A")
    t1.alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A2:K2")
    t2 = ws["A2"]
    t2.value = "सामाजिक सुरक्षा भत्ता प्राप्त गर्ने लाभग्राही व्यक्तिहरूको विस्तृत विवरण सूची"
    t2.font = Font(name="Kalimati", size=11, bold=True, color="475569")
    t2.alignment = Alignment(horizontal="center", vertical="center")
    
    headers = [
        "क्र.सं.", "लाभग्राही परिचय नं.", "लाभग्राहीको नाम", "लिङ्ग", "नागरिकता / NID नं.",
        "भत्ताको प्रकार (Category)", "वडा नं.", "टोल / गाउँ", "बैंकको नाम", "बैंक खाता नम्बर",
        "मासिक रकम (रु.)", "अवस्था"
    ]
    ws.append([])
    ws.append(headers)
    
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    for col_idx in range(1, len(headers) + 1):
        c = ws.cell(row=4, column=col_idx)
        c.font = Font(name="Kalimati", size=10, bold=True, color="FFFFFF")
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    
    row_idx = 5
    for idx, b in enumerate(beneficiaries, 1):
        row_data = [
            idx,
            b.beneficiary_id,
            b.full_name,
            b.gender or '-',
            b.citizenship_or_nid_no or '-',
            b.category,
            f"वडा नं. {b.ward_number}",
            b.tole_village or '-',
            b.bank_name or '-',
            b.bank_account_no or '-',
            float(b.monthly_allowance_npr or 0),
            b.status
        ]
        ws.append(row_data)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = Font(name="Kalimati", size=9)
            cell.border = thin_border
            if col_idx in [1, 4, 7, 12]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 11:
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            if row_idx % 2 == 0:
                cell.fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        row_idx += 1
        
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    filename = f"Social_Security_Beneficiaries_{datetime.datetime.now().strftime('%Y-%m-%d')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@router.get('/stats')
def get_social_security_stats():
    return {'total': 0}
