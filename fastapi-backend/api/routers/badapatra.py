from fastapi import APIRouter, HTTPException, Query, Body
from fastapi.responses import StreamingResponse
from typing import Optional, List, Dict, Any
import os
import json
import re
import time
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from schemas.badapatra import (
    BadapatraServiceItem,
    BadapatraDepartment,
    BadapatraOfficial,
    BadapatraSummaryResponse,
    BadapatraServicesListResponse,
    BadapatraDepartmentsResponse,
    BadapatraOfficialsResponse
)

router = APIRouter(
    prefix="/api/v1/badapatra",
    tags=["Badapatra (Citizen Charter)"]
)

DATA_CACHE: Optional[Dict[str, Any]] = None
DATA_FILE_PATH: Optional[str] = None

def parse_documents(doc_str: Any) -> List[str]:
    if isinstance(doc_str, list):
        return doc_str
    if not doc_str or not isinstance(doc_str, str):
        return ["आवश्यक कागजातको विवरण वडा/शाखाबाट उपलब्ध हुनेछ।"]
    split = [s.strip() for s in re.split(r'(?=[०-९\d]+[\)\.])', doc_str) if s.strip()]
    return split if split else [doc_str]

def normalize_item(raw: Dict[str, Any], default_type: str = "ward") -> Dict[str, Any]:
    is_ward = raw.get("isWardLevel")
    if is_ward is None:
        is_ward = (
            raw.get("type") == "ward"
            or raw.get("sourceType") == "ward"
            or default_type == "ward"
            or raw.get("category") == "वडा कार्यालय"
            or ("वडा" in str(raw.get("provider", "")))
        )

    title = raw.get("title") or raw.get("serviceName") or "सेवा"
    category = raw.get("category") or raw.get("serviceCategory") or ("वडा कार्यालय" if is_ward else "गाउँकार्यपालिका")
    provider = raw.get("provider") or raw.get("departmentName") or ("वडा कार्यालय" if is_ward else "गाउँपालिका केन्द्र")
    fee = raw.get("price") or raw.get("fee") or "निःशुल्क"
    processing_time = raw.get("serviceTime") or raw.get("processingTime") or "सोही दिन"
    official = raw.get("official") or raw.get("responsibleOfficial") or ("वडा सचिव / वडा अध्यक्ष" if is_ward else "शाखा प्रमुख")
    room = raw.get("room") or raw.get("roomNo") or ("वडा प्रशासन कक्ष" if is_ward else "प्रशासन कक्ष")
    docs = raw.get("requiredDocuments") if isinstance(raw.get("requiredDocuments"), list) else parse_documents(raw.get("document"))

    item_id = str(raw.get("id") or int(time.time() * 1000))
    raw_ward_id = raw.get("wardId")
    raw_ward_no = raw.get("wardNo")
    ward_id = raw_ward_id if raw_ward_id is not None else ("all" if is_ward else None)
    ward_no = raw_ward_no if raw_ward_no is not None else (raw_ward_id if raw_ward_id is not None else ("all" if is_ward else None))

    return {
        **raw,
        "id": item_id,
        "type": "ward" if is_ward else "palika",
        "sourceType": "ward" if is_ward else "palika",
        "isWardLevel": is_ward,
        "category": category,
        "serviceCategory": category,
        "title": title,
        "serviceName": title,
        "price": fee,
        "fee": fee,
        "serviceTime": processing_time,
        "processingTime": processing_time,
        "provider": provider,
        "departmentName": provider,
        "official": official,
        "responsibleOfficial": official,
        "room": room,
        "roomNo": room,
        "document": raw.get("document") if isinstance(raw.get("document"), str) else " ".join(docs),
        "requiredDocuments": docs,
        "remarks": raw.get("remarks") or raw.get("procedure") or "",
        "procedure": raw.get("procedure") or raw.get("remarks") or "फारम भरी आवश्यक कागजात संलग्न गरी कार्यालयमा पेश गर्ने।",
        "wardId": ward_id,
        "wardNo": ward_no,
        "isActive": raw.get("isActive", True) is not False,
    }

def get_data_bundle() -> Dict[str, Any]:
    global DATA_CACHE, DATA_FILE_PATH
    if DATA_CACHE is not None:
        return DATA_CACHE

    base_dir = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    possible_paths = [
        os.path.join(base_dir, "data", "badapatra-data.json"),
        os.path.join(base_dir, "..", "backend", "src", "badapatra", "data", "badapatra-data.json"),
        os.path.join(base_dir, "..", "frontend", "public", "badapatra-data.json"),
    ]

    for p in possible_paths:
        norm_path = os.path.normpath(p)
        if os.path.exists(norm_path):
            try:
                with open(norm_path, "r", encoding="utf-8") as f:
                    parsed = json.load(f)
                DATA_FILE_PATH = norm_path
                ward_services = [normalize_item(item, "ward") for item in parsed.get("wardServices", [])]
                palika_services = [normalize_item(item, "palika") for item in parsed.get("palikaServices", [])]
                stats = {
                    "totalServices": len(ward_services) + len(palika_services),
                    "wardServicesCount": len(ward_services),
                    "palikaServicesCount": len(palika_services),
                    "departmentsCount": len(parsed.get("departments", [])),
                    "officialsCount": len(parsed.get("officials", [])),
                }
                DATA_CACHE = {
                    **parsed,
                    "stats": stats,
                    "summary": stats,
                    "wardServices": ward_services,
                    "palikaServices": palika_services,
                    "departments": parsed.get("departments", []),
                    "officials": parsed.get("officials", []),
                }
                return DATA_CACHE
            except Exception:
                pass

    # Fallback default empty
    empty_stats = {"totalServices": 0, "wardServicesCount": 0, "palikaServicesCount": 0, "departmentsCount": 0, "officialsCount": 0}
    DATA_CACHE = {
        "municipality": "पाँचपोखरी थाङपाल गाउँपालिका",
        "sourceUrl": "https://badapatra.smartptrm.com",
        "stats": empty_stats,
        "wardServices": [],
        "palikaServices": [],
        "departments": [],
        "officials": []
    }
    return DATA_CACHE

def persist_to_disk():
    global DATA_CACHE, DATA_FILE_PATH
    if not DATA_CACHE or not DATA_FILE_PATH:
        return
    try:
        to_save = {
            "municipality": DATA_CACHE.get("municipality", "पाँचपोखरी थाङपाल गाउँपालिका"),
            "district": DATA_CACHE.get("district", "सिन्धुपाल्चोक"),
            "portalUrl": DATA_CACHE.get("portalUrl", "https://badapatra.smartptrm.com/"),
            "summary": {
                "totalServices": len(DATA_CACHE["wardServices"]) + len(DATA_CACHE["palikaServices"]),
                "wardServicesCount": len(DATA_CACHE["wardServices"]),
                "palikaServicesCount": len(DATA_CACHE["palikaServices"]),
                "departmentsCount": len(DATA_CACHE.get("departments", [])),
                "officialsCount": len(DATA_CACHE.get("officials", [])),
            },
            "wardServices": DATA_CACHE["wardServices"],
            "palikaServices": DATA_CACHE["palikaServices"],
            "departments": DATA_CACHE.get("departments", []),
            "officials": DATA_CACHE.get("officials", []),
        }
        with open(DATA_FILE_PATH, "w", encoding="utf-8") as f:
            json.dump(to_save, f, ensure_ascii=False, indent=2)
    except Exception:
        pass

@router.get("/summary", response_model=BadapatraSummaryResponse)
def get_summary():
    data = get_data_bundle()
    total = len(data["wardServices"]) + len(data["palikaServices"])
    stats = {
        "totalServices": total,
        "wardServicesCount": len(data["wardServices"]),
        "palikaServicesCount": len(data["palikaServices"]),
        "departmentsCount": len(data["departments"]),
        "officialsCount": len(data["officials"]),
    }
    return BadapatraSummaryResponse(
        success=True,
        municipality=data.get("municipality", "पाँचपोखरी थाङपाल गाउँपालिका"),
        sourceUrl=data.get("sourceUrl") or data.get("portalUrl"),
        lastSync=data.get("fetchedAt") or data.get("lastSync"),
        stats=stats
    )

@router.get("/ward-letters", response_model=BadapatraServicesListResponse)
def get_ward_letters(
    q: Optional[str] = Query(None, description="Search query"),
    wardId: Optional[int] = Query(None, description="Filter by ward number")
):
    data = get_data_bundle()
    services = data.get("wardServices", [])
    if wardId is not None and wardId > 0:
        services = [
            s for s in services
            if s.get("wardId") == wardId or s.get("wardNo") == wardId or s.get("wardId") == "all" or not s.get("wardId")
        ]
    if q and q.strip():
        term = q.strip().lower()
        services = [
            s for s in services
            if (
                term in str(s.get("serviceName", "")).lower()
                or term in str(s.get("title", "")).lower()
                or term in str(s.get("serviceCategory", "")).lower()
                or term in str(s.get("category", "")).lower()
                or term in str(s.get("responsibleOfficial", "")).lower()
                or term in str(s.get("official", "")).lower()
            )
        ]
    return BadapatraServicesListResponse(
        success=True,
        count=len(services),
        data=services,
        result=services
    )

@router.get("/palika-letters", response_model=BadapatraServicesListResponse)
def get_palika_letters(
    q: Optional[str] = Query(None, description="Search query"),
    departmentId: Optional[str] = Query(None, description="Filter by department ID")
):
    data = get_data_bundle()
    services = data.get("palikaServices", [])
    if departmentId:
        services = [
            s for s in services
            if s.get("departmentId") == departmentId or s.get("category") == departmentId
        ]
    if q and q.strip():
        term = q.strip().lower()
        services = [
            s for s in services
            if (
                term in str(s.get("serviceName", "")).lower()
                or term in str(s.get("title", "")).lower()
                or term in str(s.get("departmentName", "")).lower()
                or term in str(s.get("provider", "")).lower()
                or term in str(s.get("responsibleOfficial", "")).lower()
                or term in str(s.get("official", "")).lower()
            )
        ]
    return BadapatraServicesListResponse(
        success=True,
        count=len(services),
        data=services,
        result=services
    )

@router.get("/services", response_model=BadapatraServicesListResponse)
def get_all_services(q: Optional[str] = Query(None, description="Search query")):
    data = get_data_bundle()
    all_services = list(data.get("wardServices", [])) + list(data.get("palikaServices", []))
    if not q or not q.strip():
        return BadapatraServicesListResponse(
            success=True,
            count=len(all_services),
            data=all_services,
            result=all_services
        )
    term = q.strip().lower()
    filtered = [
        s for s in all_services
        if (
            term in str(s.get("serviceName", "")).lower()
            or term in str(s.get("title", "")).lower()
            or term in str(s.get("serviceCategory", "")).lower()
            or term in str(s.get("category", "")).lower()
            or term in str(s.get("departmentName", "")).lower()
            or term in str(s.get("provider", "")).lower()
            or term in str(s.get("responsibleOfficial", "")).lower()
            or term in str(s.get("official", "")).lower()
        )
    ]
    return BadapatraServicesListResponse(
        success=True,
        count=len(filtered),
        data=filtered,
        result=filtered
    )

@router.get("/services/{id}")
def get_service_by_id(id: str):
    data = get_data_bundle()
    all_services = list(data.get("wardServices", [])) + list(data.get("palikaServices", []))
    for s in all_services:
        if str(s.get("id")) == str(id):
            return {"success": True, "data": s, "result": s}
    raise HTTPException(status_code=404, detail=f"Service with id {id} not found in citizen charter")

@router.post("/services")
def create_service(body: Dict[str, Any] = Body(...)):
    data = get_data_bundle()
    is_ward = body.get("isWardLevel", body.get("type") == "ward" or body.get("sourceType") == "ward")
    new_id = str(int(time.time() * 1000))
    item = normalize_item({**body, "id": new_id}, "ward" if is_ward else "palika")

    if is_ward:
        data["wardServices"].insert(0, item)
    else:
        data["palikaServices"].insert(0, item)

    persist_to_disk()
    return {"success": True, "message": "नागरिक वडापत्र सेवा सफलतापूर्वक थप गरियो", "data": item, "result": item}

@router.put("/services/{id}")
def update_service(id: str, body: Dict[str, Any] = Body(...)):
    data = get_data_bundle()
    found = False
    updated_item = None

    for idx, item in enumerate(data["wardServices"]):
        if str(item.get("id")) == str(id):
            is_still_ward = body.get("isWardLevel", item.get("isWardLevel"))
            updated_item = normalize_item({**item, **body, "id": item["id"]}, "ward" if is_still_ward else "palika")
            if is_still_ward:
                data["wardServices"][idx] = updated_item
            else:
                data["wardServices"].pop(idx)
                data["palikaServices"].insert(0, updated_item)
            found = True
            break

    if not found:
        for idx, item in enumerate(data["palikaServices"]):
            if str(item.get("id")) == str(id):
                is_still_ward = body.get("isWardLevel", item.get("isWardLevel"))
                updated_item = normalize_item({**item, **body, "id": item["id"]}, "ward" if is_still_ward else "palika")
                if not is_still_ward:
                    data["palikaServices"][idx] = updated_item
                else:
                    data["palikaServices"].pop(idx)
                    data["wardServices"].insert(0, updated_item)
                found = True
                break

    if not found or not updated_item:
        raise HTTPException(status_code=404, detail=f"Service with id {id} not found to update")

    persist_to_disk()
    return {"success": True, "message": "नागरिक वडापत्र सेवा सफलतापूर्वक अद्यावधिक गरियो", "data": updated_item, "result": updated_item}

@router.delete("/services/{id}")
def delete_service(id: str):
    data = get_data_bundle()
    prev_ward = len(data["wardServices"])
    prev_palika = len(data["palikaServices"])

    data["wardServices"] = [s for s in data["wardServices"] if str(s.get("id")) != str(id)]
    data["palikaServices"] = [s for s in data["palikaServices"] if str(s.get("id")) != str(id)]

    if len(data["wardServices"]) == prev_ward and len(data["palikaServices"]) == prev_palika:
        raise HTTPException(status_code=404, detail=f"Service with id {id} not found to delete")

    persist_to_disk()
    return {"success": True, "message": "नागरिक वडापत्र सेवा सफलतापूर्वक हटाइयो", "deletedId": id}

@router.get("/departments", response_model=BadapatraDepartmentsResponse)
def get_departments(q: Optional[str] = Query(None, description="Search query")):
    data = get_data_bundle()
    departments = data.get("departments", [])
    if q and q.strip():
        term = q.strip().lower()
        departments = [
            d for d in departments
            if term in str(d.get("name", "")).lower() or term in str(d.get("headName", "")).lower()
        ]
    return BadapatraDepartmentsResponse(
        success=True,
        count=len(departments),
        data=departments,
        result=departments
    )

@router.get("/officials", response_model=BadapatraOfficialsResponse)
def get_officials(
    q: Optional[str] = Query(None, description="Search query"),
    department: Optional[str] = Query(None, description="Department filter"),
    wardNo: Optional[int] = Query(None, description="Ward number")
):
    data = get_data_bundle()
    officials = data.get("officials", [])
    if wardNo is not None:
        officials = [o for o in officials if o.get("wardNo") == wardNo]
    if department:
        dep_lower = department.lower()
        officials = [o for o in officials if dep_lower in str(o.get("departmentName", "")).lower()]
    if q and q.strip():
        term = q.strip().lower()
        officials = [
            o for o in officials
            if (
                term in str(o.get("name", "")).lower()
                or term in str(o.get("designation", "")).lower()
                or term in str(o.get("departmentName", "")).lower()
            )
        ]
    return BadapatraOfficialsResponse(
        success=True,
        count=len(officials),
        data=officials,
        result=officials
    )


def generate_badapatra_workbook() -> openpyxl.Workbook:
    data = get_data_bundle()
    services = list(data.get("wardServices", [])) + list(data.get("palikaServices", []))

    wb = openpyxl.Workbook()

    # Sheet 1: Citizen_Charter_Standard (matching exact 42 columns of CITIZEN_CHARTER template)
    ws1 = wb.active
    ws1.title = "Citizen_Charter_Standard"

    cols_42 = [
        'id', 'created_at', 'last_modified', 'deleted_at',
        'service_details', 'service_details_ne', 'service_details_en', 'service_details_ol',
        'required_process', 'required_documents', 'required_documents_ne', 'required_documents_en', 'required_documents_ol',
        'fees', 'fees_ne', 'fees_en', 'fees_ol',
        'processing_time', 'processing_time_ne', 'processing_time_en', 'processing_time_ol',
        'responsible_person_or_section', 'responsible_person_or_section_ne', 'responsible_person_or_section_en', 'responsible_person_or_section_ol',
        'flow_time', 'flow_time_ne', 'flow_time_en', 'flow_time_ol',
        'complaint_officer', 'service_code', 'service_type', 'online_service_link',
        'providing_department', 'fee_details', 'process_steps', 'related_law',
        'feature_icon', 'is_active', 'department', 'ward', 'ward_department'
    ]
    ws1.append(cols_42)

    now_str = "2026-09-24 10:55:37+00:00"
    for s in services:
        s_id = str(s.get("id", ""))
        title = s.get("title") or s.get("serviceName") or ""
        docs = s.get("requiredDocuments") or []
        docs_str = "\n".join(docs) if isinstance(docs, list) else str(s.get("document") or "")
        fee = s.get("fee") or s.get("price") or "निःशुल्क"
        time_req = s.get("processingTime") or s.get("serviceTime") or "सोही दिन"
        official = s.get("responsibleOfficial") or s.get("official") or "वडा सचिव / शाखा प्रमुख"
        is_ward = s.get("isWardLevel", True)
        dept = s.get("departmentName") or s.get("provider") or ("वडा कार्यालय" if is_ward else "गाउँपालिका केन्द्र")
        room = s.get("roomNo") or s.get("room") or ("वडा प्रशासन कक्ष" if is_ward else "प्रशासन कक्ष")
        proc = s.get("procedure") or s.get("remarks") or "फारम भरी आवश्यक कागजात संलग्न गरी कार्यालयमा पेश गर्ने।"
        ward_val = "सबै वडा (१-८)" if is_ward else "गाउँपालिका केन्द्र"
        complaint = "प्रमुख प्रशासकीय अधिकृत / गाउँपालिका उपाध्यक्ष"

        row = [
            s_id, now_str, now_str, "", title, title, "", "", proc, docs_str, docs_str, "", "",
            fee, fee, "", "", time_req, time_req, "", "", official, official, "", "", "", "", "", "",
            complaint, f"S-{s_id}", "वडा तह" if is_ward else "गाउँपालिका तह", "https://eservice.panchpokharimun.gov.np",
            dept, fee, proc, "स्थानीय सरकार सञ्चालन ऐन, २०७४", "file-text", "TRUE",
            s.get("category") or dept, ward_val, room
        ]
        ws1.append(row)

    # Sheet 2: नागरिक_वडापत्र_विस्तृत
    ws2 = wb.create_sheet(title="नागरिक_वडापत्र_विस्तृत")
    nepali_cols = [
        'क्र.सं.', 'सेवाको नाम (Service Name)', 'तह (Level)', 'शाखा / वडा (Department / Ward)',
        'दस्तुर (Fee)', 'लाग्ने समय (Processing Time)', 'जिम्मेवार अधिकारी (Responsible Official)',
        'सेवा प्रवाह कक्ष (Room No)', 'आवश्यक कागजातहरू (Required Documents)', 'सेवा प्रक्रिया (Procedure)',
        'गुनासो सुन्ने अधिकारी (Grievance Officer)'
    ]
    ws2.append(nepali_cols)
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    for col_idx in range(1, len(nepali_cols) + 1):
        cell = ws2.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for idx, s in enumerate(services, start=1):
        title = s.get("title") or s.get("serviceName") or ""
        docs = s.get("requiredDocuments") or []
        docs_str = "\n".join(docs) if isinstance(docs, list) else str(s.get("document") or "")
        fee = s.get("fee") or s.get("price") or "निःशुल्क"
        time_req = s.get("processingTime") or s.get("serviceTime") or "सोही दिन"
        official = s.get("responsibleOfficial") or s.get("official") or "वडा सचिव / शाखा प्रमुख"
        is_ward = s.get("isWardLevel", True)
        dept = s.get("departmentName") or s.get("provider") or ("वडा कार्यालय" if is_ward else "गाउँपालिका केन्द्र")
        room = s.get("roomNo") or s.get("room") or ("वडा प्रशासन कक्ष" if is_ward else "प्रशासन कक्ष")
        proc = s.get("procedure") or s.get("remarks") or "फारम भरी आवश्यक कागजात संलग्न गरी कार्यालयमा पेश गर्ने।"
        complaint = "प्रमुख प्रशासकीय अधिकृत / गाउँपालिका उपाध्यक्ष"

        ws2.append([
            idx, title, "वडा तह" if is_ward else "पालिका केन्द्र", dept, fee, time_req,
            official, room, docs_str, proc, complaint
        ])

    for col in ws2.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = col[0].column_letter
        ws2.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

    return wb


@router.get("/export/xlsx")
def export_badapatra_xlsx():
    """नागरिक वडापत्रलाई Excel (.xlsx) मा Export गर्नुहोस् (CITIZEN_CHARTER format)"""
    wb = generate_badapatra_workbook()
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    timestamp = time.strftime("%Y-%m-%d_%H_%M_%S")
    filename = f"CITIZEN_CHARTER_{timestamp}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )



@router.get('/citizens-charter')
def get_badapatra_citizens_charter():
    return []
